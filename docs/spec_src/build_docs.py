"""Builds the client approval documents from rules.json and tests.json:

  docs/WT1.1_Functional_Specification_v<ver>.docx   - rules with IDs, client OK
                                                        checkboxes and comment column
  docs/WT1.1_Test_Matrix_v<ver>.xlsx                 - test cases per rule, status
                                                        dropdown, summary, decisions

Run:  py -3 build_docs.py      (needs: pip install python-docx openpyxl)
Edit rules.json / tests.json, bump "version" there, and run again.
"""
import json
import os
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SRC = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.dirname(SRC)

with open(os.path.join(SRC, "rules.json"), encoding="utf-8") as f:
    DATA = json.load(f)
with open(os.path.join(SRC, "tests.json"), encoding="utf-8") as f:
    TESTS = json.load(f)

VER = DATA["version"]
DOCX_PATH = os.path.join(OUT, f"WT1.1_Functional_Specification_v{VER}.docx")
XLSX_PATH = os.path.join(OUT, f"WT1.1_Test_Matrix_v{VER}.xlsx")

NAVY, BLUE = "1F3864", "2F5496"
HEAD_FILL, CLIENT_FILL, BORDER = "D9E2F3", "FFF8DC", "B4C6E7"


# ============================================================ WORD ==========

def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_fill}"/>'))


def set_cell_width(cell, cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), str(int(cm * 567)))
    w.set(qn("w:type"), "dxa")
    tc_pr.append(w)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)


def cell_text(cell, text, bold=False, size=9, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def add_checkbox(paragraph, label):
    """Clickable Word checkbox (content control) followed by a label."""
    sdt = parse_xml(
        '<w:sdt %s xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">'
        '<w:sdtPr><w:rPr><w:rFonts w:ascii="MS Gothic" w:eastAsia="MS Gothic" w:hAnsi="MS Gothic"/>'
        '<w:sz w:val="18"/></w:rPr>'
        '<w14:checkbox><w14:checked w14:val="0"/>'
        '<w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
        '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/></w14:checkbox></w:sdtPr>'
        '<w:sdtContent><w:r><w:rPr><w:rFonts w:ascii="MS Gothic" w:eastAsia="MS Gothic" w:hAnsi="MS Gothic"/>'
        '<w:sz w:val="18"/></w:rPr><w:t>☐</w:t></w:r></w:sdtContent></w:sdt>' % nsdecls("w"))
    paragraph._p.append(sdt)
    run = paragraph.add_run(" " + label + "   ")
    run.font.size = Pt(9)


def make_table(doc, headers, widths_cm, client_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        cell_text(c, h, bold=True, size=9, color=NAVY)
        shade(c, HEAD_FILL)
        set_cell_width(c, widths_cm[i])
    table._client_cols = client_cols
    table._widths = widths_cm
    return table


def add_row(table, values, bold_first=True):
    row = table.add_row()
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, v in enumerate(values):
        c = row.cells[i]
        set_cell_width(c, table._widths[i])
        if v == "__CHECK__":
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_checkbox(p, "Yes")
            add_checkbox(p, "No")
        else:
            cell_text(c, v, bold=(bold_first and i == 0))
        if i in table._client_cols:
            shade(c, CLIENT_FILL)
    return row


def add_field(paragraph, instr):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    run._r.append(fld)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-1" \\h \\z \\u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "Right-click and choose Update Field to show the table of contents."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, txt, end):
        run._r.append(el)
    # ask Word to refresh fields (TOC, page numbers) when the file is opened
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
    settings.append(upd)


def build_docx():
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.8))

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    for name, size, color in (("Heading 1", 14, NAVY), ("Heading 2", 12, BLUE), ("Title", 24, NAVY)):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        fonts = st.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)          # theme font would override Arial
        fonts.set(qn("w:eastAsia"), "Arial")

    # header / footer
    hp = sec.header.paragraphs[0]
    hr = hp.add_run(f"{DATA['product']}  |  Functional Specification v{VER}")
    hr.font.size = Pt(8); hr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page "); r.font.size = Pt(8)
    add_field(fp, "PAGE")
    r = fp.add_run(" of "); r.font.size = Pt(8)
    add_field(fp, "NUMPAGES")

    # ---- cover ----
    doc.add_paragraph().paragraph_format.space_after = Pt(40)
    doc.add_paragraph("Functional Specification", style="Title")
    sub = doc.add_paragraph()
    sr = sub.add_run(f"{DATA['product']} – Operating Modes and Protections")
    sr.font.size = Pt(13); sr.font.color.rgb = RGBColor.from_string(BLUE)

    t = make_table(doc, ["Item", "Details"], [5.0, 12.4])
    for k, v in (("Document version", VER), ("Date", DATA["date"]),
                 ("Firmware reviewed", DATA["firmware"]),
                 ("Status", "Draft for client approval"),
                 ("Companion document", f"WT1.1_Test_Matrix_v{VER}.xlsx (test cases for every rule)")):
        add_row(t, [k, v])

    doc.add_heading("How to review this document", level=2)
    for line in (
        "Every rule has an ID (for example AUTO-04). Please tick Yes or No in the “Client OK?” "
        "column for each rule, and write any change in the “Client Comment” column (yellow cells).",
        "Answer the open questions in the “Open points” section.",
        "You can also select any text and use Word’s Review → New Comment.",
        "Sign the “Client approval” section at the end. Any later change becomes a new version of this document.",
        "When reporting a problem after approval, please quote the rule ID.",
    ):
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Contents", level=2)
    add_toc(doc)
    doc.add_page_break()

    # ---- terms ----
    doc.add_heading("1. Terms used", level=1)
    t = make_table(doc, ["Term", "Meaning"], [4.2, 13.2])
    for term, meaning in DATA["terms"]:
        add_row(t, [term, meaning])

    # ---- modules ----
    n = 2
    for m in DATA["modules"]:
        doc.add_heading(f"{n}. {m['name']}", level=1)
        if m.get("note"):
            p = doc.add_paragraph(m["note"])
            p.runs[0].italic = True
        t = make_table(doc, ["ID", "Rule", "Client OK?", "Client Comment"],
                       [1.9, 9.1, 2.6, 3.8], client_cols=(2, 3))
        for rid, text in m["rules"]:
            add_row(t, [rid, text, "__CHECK__", ""])
        n += 1

    # ---- decisions ----
    doc.add_heading(f"{n}. Open points – client decision needed", level=1)
    doc.add_paragraph("These behaviours are implemented as described under “Current behaviour”. "
                      "Please choose an option or describe what is required.")
    t = make_table(doc, ["ID", "Topic", "Current behaviour", "Options", "Client decision / comment"],
                   [1.6, 2.6, 5.4, 3.9, 3.9], client_cols=(4,))
    for d in DATA["decisions"]:
        add_row(t, [d[0], d[1], d[2], d[3], ""])
    n += 1

    # ---- revision history ----
    doc.add_heading(f"{n}. Revision history", level=1)
    t = make_table(doc, ["Version", "Date", "Changes", "By"], [2.0, 2.6, 9.8, 3.0])
    for h in DATA.get("history", []):
        add_row(t, h, bold_first=False)
    add_row(t, ["", "", "", ""], bold_first=False)
    n += 1

    # ---- approval ----
    doc.add_heading(f"{n}. Client approval", level=1)
    doc.add_paragraph("By signing, the client confirms that the rules in this document, together with the "
                      "comments and decisions written in it, describe the required behaviour.")
    t = make_table(doc, ["Item", "Client"], [5.0, 12.4], client_cols=(1,))
    for item in ("Overall comments", "Approved (Yes / Yes with changes / No)", "Name",
                 "Designation / Company", "Signature", "Date"):
        row = add_row(t, [item, ""], bold_first=False)
        if item == "Overall comments":
            row.height = Cm(3.0)
        elif item == "Signature":
            row.height = Cm(1.6)

    doc.save(DOCX_PATH)
    return DOCX_PATH


# ============================================================ EXCEL =========

ARIAL = "Arial"
thin = Side(style="thin", color="B4C6E7")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
HEAD_BG = PatternFill("solid", fgColor=NAVY)
CLIENT_BG = PatternFill("solid", fgColor=CLIENT_FILL)
BODY = Font(name=ARIAL, size=9)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def write_header(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_BG, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def body_cell(ws, row, col, value, client=False, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=ARIAL, size=9, bold=bold)
    c.alignment = WRAP_TOP
    c.border = BOX
    if client:
        c.fill = CLIENT_BG
    return c


def build_xlsx():
    wb = Workbook()
    modules = [(m["id"], m["name"]) for m in DATA["modules"]]
    module_name = dict(modules)

    # ---------------- Instructions ----------------
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 95
    ws["A1"] = f"{DATA['product']} – Test Matrix v{VER}"
    ws["A1"].font = Font(name=ARIAL, bold=True, size=14, color=NAVY)
    ws["A2"] = f"Date {DATA['date']}. Firmware: {DATA['firmware']}. " \
               f"Rules are described in WT1.1_Functional_Specification_v{VER}.docx."
    ws["A2"].font = BODY
    lines = [
        ("Sheet", "Purpose"),
        ("Test Cases", "One row per test. Each test is linked to a rule ID of the specification."),
        ("Summary", "Pass / Fail count per module (calculated automatically from the Status column)."),
        ("Rules", "All rules with a Client OK? dropdown and a Client Comment column."),
        ("Decisions", "Open points that need a client decision."),
        ("", ""),
        ("Columns to fill (yellow)", "Actual Result, Status, Tested By, Date, Client Comment. "
                                     "Developer Remark is filled by the development team."),
        ("Status values", "Not Run = not tested yet   Pass = result matches Expected   "
                          "Fail = result differs (describe it in Actual Result)   "
                          "Blocked = could not be tested (say why in Actual Result)"),
        ("New problem found?", "Add a new row at the bottom with the next Test ID and the rule ID it belongs "
                               "to (or 'NEW' if no rule covers it). A 'NEW' row means the specification "
                               "needs a new rule."),
    ]
    r = 4
    for a, b in lines:
        ca, cb = ws.cell(row=r, column=1, value=a), ws.cell(row=r, column=2, value=b)
        for c in (ca, cb):
            c.font = Font(name=ARIAL, size=10, bold=(r == 4)); c.alignment = WRAP_TOP
        if r == 4:
            ca.fill = cb.fill = PatternFill("solid", fgColor=HEAD_FILL)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Example of a filled row").font = Font(name=ARIAL, bold=True, size=10, color=NAVY)
    r += 1
    example = [("Test ID", "TC-AUTO-07"), ("Rule ID", "AUTO-04"),
               ("Precondition", "Auto stopped at 100%"), ("Steps / Event", "Lower level to 25%"),
               ("Expected Result", "Motor ON"), ("Actual Result", "Motor ON after about 6 s"),
               ("Status", "Pass"), ("Tested By", "R. Kumar"), ("Date", "27-09-2026"),
               ("Client Comment", "OK"), ("Developer Remark", "")]
    for a, b in example:
        ws.cell(row=r, column=1, value=a).font = BODY
        c = ws.cell(row=r, column=2, value=b); c.font = BODY
        r += 1

    # ---------------- Rules ----------------
    wr = wb.create_sheet("Rules")
    headers = ["Rule ID", "Module", "Rule", "Client OK?", "Client Comment"]
    write_header(wr, 1, headers, [11, 16, 80, 12, 40])
    dv_ok = DataValidation(type="list", formula1='"Yes,No,Change"', allow_blank=True)
    wr.add_data_validation(dv_ok)
    row = 2
    for m in DATA["modules"]:
        for rid, text in m["rules"]:
            body_cell(wr, row, 1, rid, bold=True)
            body_cell(wr, row, 2, m["name"])
            body_cell(wr, row, 3, text)
            body_cell(wr, row, 4, None, client=True)
            body_cell(wr, row, 5, None, client=True)
            dv_ok.add(f"D{row}")
            row += 1
    last_rule_row = row - 1
    wr.freeze_panes = "A2"
    wr.auto_filter.ref = f"A1:E{last_rule_row}"

    # ---------------- Test Cases ----------------
    wt = wb.create_sheet("Test Cases")
    headers = ["Test ID", "Module", "Rule ID", "Rule (from Rules sheet)", "Precondition / setup",
               "Steps / event", "Expected result", "Actual result", "Status", "Tested by",
               "Date", "Client Comment", "Developer Remark"]
    widths = [12, 13, 10, 36, 28, 30, 32, 28, 11, 12, 11, 28, 24]
    write_header(wt, 1, headers, widths)
    dv_status = DataValidation(type="list", formula1='"Not Run,Pass,Fail,Blocked"', allow_blank=False)
    wt.add_data_validation(dv_status)

    counters = {}
    row = 2
    for mod, rule, pre, steps, expected in TESTS:
        counters[mod] = counters.get(mod, 0) + 1
        tid = f"TC-{mod}-{counters[mod]:02d}"
        body_cell(wt, row, 1, tid, bold=True)
        body_cell(wt, row, 2, module_name[mod])
        body_cell(wt, row, 3, rule)
        body_cell(wt, row, 4, f'=IFERROR(INDEX(Rules!$C$2:$C${last_rule_row},'
                              f'MATCH(C{row},Rules!$A$2:$A${last_rule_row},0)),"")')
        body_cell(wt, row, 5, pre)
        body_cell(wt, row, 6, steps)
        body_cell(wt, row, 7, expected)
        body_cell(wt, row, 8, None, client=True)
        body_cell(wt, row, 9, "Not Run", client=True)
        body_cell(wt, row, 10, None, client=True)
        body_cell(wt, row, 11, None, client=True)
        body_cell(wt, row, 12, None, client=True)
        body_cell(wt, row, 13, None)
        row += 1
    last_test_row = row - 1
    dv_status.add(f"I2:I{last_test_row + 200}")   # also covers rows added later
    wt.freeze_panes = "B2"
    wt.auto_filter.ref = f"A1:M{last_test_row}"
    status_rng = f"I2:I{last_test_row + 200}"
    for text, color in (("Pass", "C6EFCE"), ("Fail", "FFC7CE"), ("Blocked", "FFEB9C")):
        wt.conditional_formatting.add(status_rng, CellIsRule(
            operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", fgColor=color)))
    wt.page_setup.orientation = "landscape"
    wt.page_setup.fitToWidth = 1
    wt.page_setup.fitToHeight = 0
    wt.sheet_properties.pageSetUpPr.fitToPage = True
    wt.print_title_rows = "1:1"

    # ---------------- Summary ----------------
    ws2 = wb.create_sheet("Summary", 1)
    ws2["A1"] = "Test summary"
    ws2["A1"].font = Font(name=ARIAL, bold=True, size=14, color=NAVY)
    ws2["A2"] = "Calculated from the Status column of the Test Cases sheet."
    ws2["A2"].font = BODY
    headers = ["Module", "Total", "Pass", "Fail", "Blocked", "Not Run", "% Pass"]
    write_header(ws2, 4, headers, [26, 9, 9, 9, 10, 10, 10])
    tc = "'Test Cases'"
    rng_mod, rng_st = f"{tc}!$B$2:$B${last_test_row + 200}", f"{tc}!$I$2:$I${last_test_row + 200}"
    r = 5
    for mid, name in modules:
        body_cell(ws2, r, 1, name, bold=True)
        body_cell(ws2, r, 2, f'=COUNTIF({rng_mod},A{r})')
        for col, st in ((3, "Pass"), (4, "Fail"), (5, "Blocked"), (6, "Not Run")):
            body_cell(ws2, r, col, f'=COUNTIFS({rng_mod},A{r},{rng_st},"{st}")')
        c = body_cell(ws2, r, 7, f'=IF(B{r}=0,0,C{r}/B{r})')
        c.number_format = "0%"
        r += 1
    body_cell(ws2, r, 1, "Total", bold=True)
    for col in range(2, 7):
        L = get_column_letter(col)
        c = body_cell(ws2, r, col, f"=SUM({L}5:{L}{r - 1})", bold=True)
    c = body_cell(ws2, r, 7, f'=IF(B{r}=0,0,C{r}/B{r})', bold=True)
    c.number_format = "0%"
    for col in range(1, 8):
        ws2.cell(row=r, column=col).fill = PatternFill("solid", fgColor=HEAD_FILL)

    # ---------------- Decisions ----------------
    wd = wb.create_sheet("Decisions")
    headers = ["ID", "Topic", "Current behaviour", "Options", "Client decision", "Client Comment"]
    write_header(wd, 1, headers, [9, 22, 55, 40, 20, 40])
    for i, d in enumerate(DATA["decisions"], start=2):
        body_cell(wd, i, 1, d[0], bold=True)
        body_cell(wd, i, 2, d[1])
        body_cell(wd, i, 3, d[2])
        body_cell(wd, i, 4, d[3])
        body_cell(wd, i, 5, None, client=True)
        body_cell(wd, i, 6, None, client=True)
    wd.freeze_panes = "A2"

    wb.calculation.fullCalcOnLoad = True
    wb.save(XLSX_PATH)
    return XLSX_PATH, last_test_row - 1


if __name__ == "__main__":
    print("Created:", build_docx())
    path, n = build_xlsx()
    print(f"Created: {path}  ({n} test cases)")
