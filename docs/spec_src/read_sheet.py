"""Reads the shared Google Sheet (link sharing: anyone with the link) and
prints what needs action: failed / blocked tests, client comments, rules
not approved and client decisions.

Run:  py -3 read_sheet.py [sheet_id]
"""
import io
import sys
import urllib.request

import openpyxl

SHEET_ID = sys.argv[1] if len(sys.argv) > 1 else "1o_aCLnH3yK23VnSosBDFFzMVE2BSKekWuaYVxLJydS8"
URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

data = urllib.request.urlopen(URL, timeout=60).read()
wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)


def rows(sheet):
    ws = wb[sheet]
    head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in r):
            yield dict(zip(head, r))


def txt(v):
    return "" if v is None else str(v).strip()


print("=== Summary ===")
ws = wb["Summary"]
for r in ws.iter_rows(min_row=4, values_only=True):
    if r[0]:
        print(" | ".join(txt(v) for v in r[:7]))

print("\n=== Tests that need action (Fail / Blocked / with comments) ===")
for t in rows("Test Cases"):
    status = txt(t.get("Status"))
    comment = txt(t.get("Client Comment"))
    if status in ("Fail", "Blocked") or comment:
        print(f"{txt(t.get('Test ID'))} [{status}] rule {txt(t.get('Rule ID'))}")
        print(f"   Steps:    {txt(t.get('Steps / event'))}")
        print(f"   Expected: {txt(t.get('Expected result'))}")
        print(f"   Actual:   {txt(t.get('Actual result'))}")
        if comment:
            print(f"   Client:   {comment}")
        if txt(t.get("Tested by")):
            print(f"   By {txt(t.get('Tested by'))} on {txt(t.get('Date'))}")

print("\n=== Rules not approved / with comments ===")
for r in rows("Rules"):
    ok, comment = txt(r.get("Client OK?")), txt(r.get("Client Comment"))
    if ok in ("No", "Change") or comment:
        print(f"{txt(r.get('Rule ID'))} [{ok or '-'}] {comment}")

print("\n=== Client decisions ===")
for d in rows("Decisions"):
    dec, comment = txt(d.get("Client decision")), txt(d.get("Client Comment"))
    print(f"{txt(d.get('ID'))} {txt(d.get('Topic'))}: {dec or '(open)'} {comment}")
